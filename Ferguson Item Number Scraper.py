"""
Ferguson Item Number Scraper
-----------------------------------
Uses a dedicated Chrome profile for a real Ferguson login session, looks up the
Ferguson item number from the search-results page, then fetches account-specific
pricing from the tile-pricing endpoint.

Requirements:
    pip install openpyxl selenium requests

Usage:
    python "Ferguson Item Number Scraper.py"
"""

import json
import os
import re
import sys
import time
from html import unescape
from urllib.parse import quote_plus

import openpyxl
import requests
import shutil
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait

# -----------------------------------------------------------------------------
# Paths and workbook settings
# -----------------------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
INPUT_FILE = os.path.join(BASE_DIR, "Ferguson Item Number Scraper Template.xlsx")
OUTPUT_FILE = os.path.join(BASE_DIR, "Ferguson Item Number Scraper Final.xlsx")
PROFILE_DIR = os.path.join(BASE_DIR, "Ferguson Item Number Scraper Profile")
DEBUG_DIR = os.path.join(BASE_DIR, "debugging")

SHEET_NAME = "Scraper Data"
MODEL_COL = 1
ITEM_COL = 3
COST_COL = 5
START_ROW = 2
SAVE_INTERVAL = 25
DEFAULT_TIMEOUT = 30

# -----------------------------------------------------------------------------
# Ferguson endpoints
# -----------------------------------------------------------------------------
FERGUSON_HOME_URL = "https://www.ferguson.com/"
FERGUSON_LOGIN_URL = "https://www.ferguson.com/s/login/"
FERGUSON_SEARCH_URL = "https://www.ferguson.com/search?q={query}&lang=en_US"
FERGUSON_PRICE_URL = (
    "https://www.ferguson.com/on/demandware.store/"
    "Sites-Ferguson-Site/en_US/Search-GetTilePricing"
)

# Reads the auth/runtime data Ferguson exposes on the page.
JS_READ_AUTH_STATE = r"""
function pickDataLayerAuth() {
    const dl = Array.isArray(window.dataLayer) ? window.dataLayer : [];
    for (let i = dl.length - 1; i >= 0; i--) {
        const item = dl[i];
        if (item && typeof item === 'object' && (
            Object.prototype.hasOwnProperty.call(item, 'login_status') ||
            Object.prototype.hasOwnProperty.call(item, 'customer_role') ||
            Object.prototype.hasOwnProperty.call(item, 'customer_type') ||
            Object.prototype.hasOwnProperty.call(item, 'user_id') ||
            Object.prototype.hasOwnProperty.call(item, 'branch_id') ||
            Object.prototype.hasOwnProperty.call(item, 'warehouse_location_id')
        )) {
            return item;
        }
    }
    return null;
}

const bodyText = (document.body && document.body.innerText ? document.body.innerText : '').toLowerCase();
return {
    url: location.href,
    title: document.title,
    body_has_sign_out: bodyText.includes('sign out') || bodyText.includes('log out') || bodyText.includes('logout'),
    data_layer_auth: pickDataLayerAuth()
};
"""


# -----------------------------------------------------------------------------
# Utility helpers
# -----------------------------------------------------------------------------
def reset_debug_dir():
    os.makedirs(DEBUG_DIR, exist_ok=True)

    for name in os.listdir(DEBUG_DIR):
        path = os.path.join(DEBUG_DIR, name)
        try:
            if os.path.isfile(path) or os.path.islink(path):
                os.remove(path)
            elif os.path.isdir(path):
                shutil.rmtree(path)
        except PermissionError:
            print(f"WARNING: Could not delete debug file or folder in use: {path}")
        except Exception as e:
            print(f"WARNING: Could not delete debug file or folder: {path} ({e})")


def save_json_debug(label, data):
    os.makedirs(DEBUG_DIR, exist_ok=True)
    safe_label = re.sub(r"[^A-Za-z0-9_.-]+", "_", label)
    path = os.path.join(DEBUG_DIR, f"debug_{safe_label}.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    return path


def save_text_debug(label, content, ext="html"):
    os.makedirs(DEBUG_DIR, exist_ok=True)
    safe_label = re.sub(r"[^A-Za-z0-9_.-]+", "_", label)
    safe_ext = re.sub(r"[^A-Za-z0-9]+", "", str(ext or "txt")) or "txt"
    path = os.path.join(DEBUG_DIR, f"debug_{safe_label}.{safe_ext}")
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    return path


def save_page_debug(driver, label):
    os.makedirs(DEBUG_DIR, exist_ok=True)
    safe_label = re.sub(r"[^A-Za-z0-9_.-]+", "_", label)
    screenshot_path = os.path.join(DEBUG_DIR, f"debug_{safe_label}.png")
    html_path = os.path.join(DEBUG_DIR, f"debug_{safe_label}.html")

    try:
        driver.save_screenshot(screenshot_path)
    except Exception:
        pass

    try:
        with open(html_path, "w", encoding="utf-8") as f:
            f.write(driver.page_source)
    except Exception:
        pass


def wait_for_dom_ready(driver, timeout=DEFAULT_TIMEOUT):
    WebDriverWait(driver, timeout).until(
        lambda d: d.execute_script("return document.readyState") in ("interactive", "complete")
    )


def dismiss_cookie_banner(driver):
    try:
        button = driver.execute_script(
            """
            const buttons = Array.from(document.querySelectorAll('button'));
            return buttons.find(b => (b.innerText || '').trim().toLowerCase() === 'dismiss') || null;
            """
        )
        if button:
            driver.execute_script("arguments[0].click();", button)
            time.sleep(0.5)
    except Exception:
        pass


def sanitize_model_number(value):
    if value is None:
        return ""

    text = str(value)

    replacements = {
        "\u2018": "",
        "\u2019": "",
        "\u201c": "",
        "\u201d": "",
        "'": "",
        '"': "",
        "\u00a0": " ",
        "–": "-",
        "—": "-",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    text = re.sub(r"\s+", " ", text).strip()
    return text


# -----------------------------------------------------------------------------
# Browser / session helpers
# -----------------------------------------------------------------------------
def create_driver():
    os.makedirs(PROFILE_DIR, exist_ok=True)
    options = Options()
    options.add_argument(f"--user-data-dir={PROFILE_DIR}")
    options.add_argument("--start-maximized")
    options.add_argument("--window-size=1920,1080")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-gpu")
    options.add_experimental_option(
        "prefs",
        {
            "profile.cookie_controls_mode": 0,
            "credentials_enable_service": False,
            "profile.password_manager_enabled": False,
        },
    )
    driver = webdriver.Chrome(options=options)
    driver.set_page_load_timeout(60)
    return driver


def get_auth_state(driver):
    return driver.execute_script(JS_READ_AUTH_STATE)


def is_authenticated_state(state):
    auth = state.get("data_layer_auth") or {}
    if auth.get("login_status") is True:
        return True
    if auth.get("customer_type") and str(auth.get("customer_type")).lower() != "guest":
        return True
    if auth.get("customer_role") and str(auth.get("customer_role")).lower() != "guest":
        return True
    if auth.get("user_id"):
        return True
    return bool(state.get("body_has_sign_out"))


def ensure_authenticated_session(driver):
    print("Opening Ferguson home page...")
    driver.get(FERGUSON_HOME_URL)
    wait_for_dom_ready(driver)
    time.sleep(2)
    dismiss_cookie_banner(driver)

    state = get_auth_state(driver)
    save_json_debug("startup_auth_state", state)
    if is_authenticated_state(state):
        print("✓ Existing authenticated session found in Chrome profile.")
        return True

    print("\nManual login required.")
    print(f"Chrome profile folder: {PROFILE_DIR}")
    print("1. Log into Ferguson in the opened browser")
    print("2. Confirm you are fully logged in")
    print("3. Return here and press Enter")

    driver.get(FERGUSON_LOGIN_URL)
    wait_for_dom_ready(driver)
    time.sleep(2)
    dismiss_cookie_banner(driver)
    save_page_debug(driver, "manual_login_prompt")

    input("Press Enter after finishing manual login... ")

    driver.get(FERGUSON_HOME_URL)
    wait_for_dom_ready(driver)
    time.sleep(2)
    dismiss_cookie_banner(driver)
    state = get_auth_state(driver)
    save_json_debug("manual_login_auth_state", state)

    if is_authenticated_state(state):
        print("✓ Manual login verified.")
        return True

    print("✗ Browser still does not appear authenticated.")
    save_page_debug(driver, "manual_login_not_authenticated")
    return False


def build_authenticated_session(driver):
    session = requests.Session()
    try:
        user_agent = driver.execute_script("return navigator.userAgent")
    except Exception:
        user_agent = "Mozilla/5.0"

    session.headers.update(
        {
            "User-Agent": user_agent,
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
            "Referer": FERGUSON_HOME_URL,
            "Origin": "https://www.ferguson.com",
        }
    )

    for cookie in driver.get_cookies():
        try:
            session.cookies.set(
                cookie["name"],
                cookie["value"],
                domain=cookie.get("domain"),
                path=cookie.get("path", "/"),
            )
        except Exception:
            continue

    return session



# -----------------------------------------------------------------------------
# Runtime context and lookup helpers
# -----------------------------------------------------------------------------
def extract_runtime_context_from_html(text):
    patterns = {
        "branch_id": r'"branch_id"\s*:\s*"([A-Za-z0-9_-]+)"',
        "warehouse_location_id": r'"warehouse_location_id"\s*:\s*"?(\d+)"?',
        "customer_main_account_number": r'"customer_main_account_number"\s*:\s*"?(\d+)"?',
    }
    context = {}
    for key, pattern in patterns.items():
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            context[key] = match.group(1)
    return context


def extract_runtime_context(driver):
    auth = (get_auth_state(driver).get("data_layer_auth") or {})
    context = {
        "branch_id": str(auth.get("branch_id") or auth.get("branchId") or ""),
        "warehouse_location_id": str(auth.get("warehouse_location_id") or auth.get("warehouseLocationId") or ""),
        "customer_main_account_number": str(auth.get("customer_main_account_number") or auth.get("customerId") or ""),
    }

    if not all(context.values()):
        html_context = extract_runtime_context_from_html(driver.page_source)
        for key, value in html_context.items():
            if not context.get(key):
                context[key] = value

    return context


def search_page_request(session, model_number):
    url = FERGUSON_SEARCH_URL.format(query=quote_plus(model_number))
    response = session.get(url, timeout=30, allow_redirects=True)
    response.raise_for_status()
    return {"url": response.url, "text": response.text}


def search_page_selenium(driver, model_number):
    url = FERGUSON_SEARCH_URL.format(query=quote_plus(model_number))
    driver.get(url)
    wait_for_dom_ready(driver)
    time.sleep(2)
    dismiss_cookie_banner(driver)
    return {"url": driver.current_url, "text": driver.page_source}


def pick_item_from_text(text, model_number):
    text_unescaped = unescape(text)
    model_clean = sanitize_model_number(model_number)

    strong_patterns = [
        rf"Item\s*(?:&#35;|#)\s*(\d{{5,10}}).{{0,500}}?Mfr\.?\s*Part\s*(?:&#35;|#)\s*{re.escape(model_clean)}",
        rf"Mfr\.?\s*Part\s*(?:&#35;|#)\s*{re.escape(model_clean)}.{{0,500}}?Item\s*(?:&#35;|#)\s*(\d{{5,10}})",
        rf"preselectedVariant=(\d{{5,10}}).{{0,500}}?{re.escape(model_clean)}",
    ]
    for pattern in strong_patterns:
        match = re.search(pattern, text_unescaped, flags=re.IGNORECASE | re.DOTALL)
        if match:
            return match.group(1)

    for match in re.finditer(re.escape(model_clean), text_unescaped, flags=re.IGNORECASE):
        start = max(0, match.start() - 800)
        end = min(len(text_unescaped), match.end() + 1500)
        chunk = text_unescaped[start:end]
        for pattern in [
            r"Item\s*(?:&#35;|#)\s*(\d{5,10})",
            r"preselectedVariant=(\d{5,10})",
            r'id=["\']p-(\d{5,10})["\']',
            r'/product/[^"\']*/(\d{5,10})\.html',
        ]:
            nearby = re.search(pattern, chunk, flags=re.IGNORECASE | re.DOTALL)
            if nearby:
                return nearby.group(1)

    # Generic fallbacks.
    for pattern in [
        r"Item\s*(?:&#35;|#)\s*(\d{5,10})",
        r"preselectedVariant=(\d{5,10})",
        r'id=["\']p-(\d{5,10})["\']',
        r'/product/[^"\']*/(\d{5,10})\.html',
        r'"productId"\s*:\s*"?(\d{5,10})"?',
        r'"productID"\s*:\s*"?(\d{5,10})"?',
        r'"sku"\s*:\s*"?(\d{5,10})"?',
    ]:
        match = re.search(pattern, text_unescaped, flags=re.IGNORECASE | re.DOTALL)
        if match:
            return match.group(1)

    return None


def get_tile_pricing(session, item_number, context, referer_url):
    params = {
        "productIDs": str(item_number),
        "shipWhseId": str(context["warehouse_location_id"]),
        "branchId": str(context["branch_id"]).upper(),
        "customerId": str(context["customer_main_account_number"]),
    }

    response = session.get(
        FERGUSON_PRICE_URL,
        params=params,
        headers={
            "Accept": "*/*",
            "Referer": referer_url,
            "X-Requested-With": "XMLHttpRequest",
        },
        timeout=30,
    )
    response.raise_for_status()
    data = response.json()

    for row in data.get("productPricing", []) or []:
        if str(row.get("id")) == str(item_number):
            raw_price = str(row.get("price", "")).replace("$", "").replace(",", "").strip()
            return float(raw_price) if raw_price else None

    return None


def lookup_item_and_price(session, driver, model_number, runtime_context):
    original_model = str(model_number).strip()
    clean_model = sanitize_model_number(original_model)
    search_terms = [term for term in [clean_model, original_model] if term]

    page = None
    for index, term in enumerate(dict.fromkeys(search_terms)):  # preserve order, remove duplicates
        page = search_page_request(session, term)
        item_number = pick_item_from_text(page["text"], term)
        if item_number:
            return {
                "item_number": item_number,
                "price": get_tile_pricing(session, item_number, runtime_context, page["url"]),
                "source": "requests",
            }

        if index == 0:
            page = search_page_selenium(driver, term)
            item_number = pick_item_from_text(page["text"], term)
            if item_number:
                return {
                    "item_number": item_number,
                    "price": get_tile_pricing(session, item_number, runtime_context, page["url"]),
                    "source": "selenium",
                }

    return {"item_number": None, "price": None, "html": page["text"] if page else None}


# -----------------------------------------------------------------------------
# Main
# -----------------------------------------------------------------------------
def main():
    reset_debug_dir()
    print(f"Script folder: {BASE_DIR}")
    print(f"Dedicated Chrome profile: {PROFILE_DIR}")

    if not os.path.exists(INPUT_FILE):
        print(f"ERROR: Input workbook not found: {INPUT_FILE}")
        sys.exit(1)

    workbook = openpyxl.load_workbook(INPUT_FILE)
    worksheet = workbook[SHEET_NAME]

    models = []
    for row in range(START_ROW, worksheet.max_row + 1):
        value = worksheet.cell(row=row, column=MODEL_COL).value
        if value and str(value).strip():
            models.append((row, str(value).strip()))

    print(f"Found {len(models)} model numbers to look up")
    print("=" * 60)

    driver = create_driver()
    try:
        if not ensure_authenticated_session(driver):
            sys.exit(1)

        session = build_authenticated_session(driver)
        runtime_context = extract_runtime_context(driver)
        save_json_debug("runtime_pricing_context", runtime_context)

        if not all(runtime_context.values()):
            print("ERROR: Missing Ferguson runtime pricing context.")
            print(json.dumps(runtime_context, indent=2))
            sys.exit(1)

        found = 0
        not_found = 0
        debug_misses_saved = 0

        for index, (row, model) in enumerate(models, 1):
            existing_item = worksheet.cell(row=row, column=ITEM_COL).value
            if existing_item:
                print(f"[{index}/{len(models)}] Row {row}: {model} — already has Item #{existing_item}, skipping")
                found += 1
                continue

            print(f"[{index}/{len(models)}] Row {row}: {model} ... ", end="", flush=True)

            try:
                result = lookup_item_and_price(session, driver, model, runtime_context)
            except Exception as exc:
                print(f"✗ {exc}")
                not_found += 1
                continue

            if result.get("item_number"):
                worksheet.cell(row=row, column=ITEM_COL).value = result["item_number"]
                if result.get("price") is not None:
                    worksheet.cell(row=row, column=COST_COL).value = result["price"]
                print(f"✓ Item #{result['item_number']}  ${result.get('price', 'N/A')} ({result.get('source')})")
                found += 1
            else:
                print("– not found")
                not_found += 1
                if result.get("html") and debug_misses_saved < 5:
                    save_text_debug(f"searchpage_{sanitize_model_number(model)}", result["html"])
                    save_page_debug(driver, f"searchpage_{sanitize_model_number(model)}")
                    debug_misses_saved += 1

            if index % SAVE_INTERVAL == 0:
                workbook.save(OUTPUT_FILE)
                print(f"  [Auto-saved at row {row}]")

        workbook.save(OUTPUT_FILE)
        print("=" * 60)
        print(f"Done! Found: {found} | Not found: {not_found}")
        print(f"Saved to: {OUTPUT_FILE}")

    finally:
        driver.quit()
        print("Browser closed.")


if __name__ == "__main__":
    main()
